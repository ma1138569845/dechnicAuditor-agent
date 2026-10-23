import base64
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT.parents[1] / "energy-audit-report" / "scripts" / "build_energy_audit_docx.py"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


BUILD_MOD = load_module("build_energy_audit_docx_ch6_layout", BUILD)


class Chapter6PhotoLayoutTests(unittest.TestCase):
    def _render(self, root: Path, captions):
        chapter_dir = root / "chapter_md"
        chapter_dir.mkdir(parents=True)
        images = {}
        manifest_images = []
        md_lines = ["# 第6章", "", "## 6.1 用电系统运行分析", ""]
        for index, caption in enumerate(captions, start=1):
            relative = f"data/images/ch6_{index}.png"
            image_path = root / relative
            image_path.parent.mkdir(parents=True, exist_ok=True)
            image_path.write_bytes(PNG)
            images[caption] = [relative]
            manifest_images.append({"caption": caption, "src": relative})
            md_lines.extend([caption, ""])
        md_path = chapter_dir / "ch6.md"
        md_path.write_text("\n".join(md_lines), encoding="utf-8")
        (root / "report_images.json").write_text(
            json.dumps({"images": manifest_images}, ensure_ascii=False),
            encoding="utf-8",
        )
        doc = Document()
        BUILD_MOD._render_md(
            doc,
            str(md_path),
            {"images": images, "project_dir": str(root)},
        )
        return doc

    @staticmethod
    def _picture_widths_cm(doc):
        return [shape.width.cm for shape in doc.inline_shapes]

    def test_two_photos_share_a_borderless_row_and_third_is_centered_alone(self):
        with tempfile.TemporaryDirectory() as td:
            doc = self._render(
                Path(td),
                ["图6.1 冷水机组", "图6.2 冷却塔", "图6.3 循环水泵"],
            )
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(len(doc.tables[0].columns), 2)
            for column in doc.tables[0].columns:
                self.assertAlmostEqual(column.width.cm, 7.3, places=1)
            self.assertEqual(doc.tables[0].cell(0, 0).text.strip(), "图6.1 冷水机组")
            self.assertEqual(doc.tables[0].cell(0, 1).text.strip(), "图6.2 冷却塔")
            borders = doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
            self.assertIsNotNone(borders)
            self.assertTrue(
                all(
                    edge.get(qn("w:val")) == "nil"
                    for edge in borders
                )
            )
            tbl_pr_children = list(doc.tables[0]._tbl.tblPr)
            borders_index = tbl_pr_children.index(borders)
            layout = doc.tables[0]._tbl.tblPr.find(qn("w:tblLayout"))
            self.assertIsNotNone(layout)
            self.assertLess(borders_index, tbl_pr_children.index(layout))
            row = doc.tables[0].rows[0]
            self.assertIsNotNone(row._tr.trPr.find(qn("w:cantSplit")))
            self.assertEqual(len(self._picture_widths_cm(doc)), 3)
            for width in self._picture_widths_cm(doc):
                self.assertAlmostEqual(width, 6.8, places=1)
            top_level_text = [p.text for p in doc.paragraphs]
            self.assertIn("图6.3 循环水泵", top_level_text)

    def test_single_chapter6_photo_is_centered_without_a_table(self):
        with tempfile.TemporaryDirectory() as td:
            doc = self._render(Path(td), ["图6.1 冷水机组"])
            self.assertEqual(len(doc.tables), 0)
            self.assertAlmostEqual(self._picture_widths_cm(doc)[0], 6.8, places=1)
            caption_para = next(p for p in doc.paragraphs if p.text == "图6.1 冷水机组")
            self.assertEqual(caption_para.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_non_chapter6_photos_keep_existing_12cm_layout(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            chapter_dir = root / "chapter_md"
            chapter_dir.mkdir()
            image_path = root / "data" / "images" / "flow.png"
            image_path.parent.mkdir(parents=True)
            image_path.write_bytes(PNG)
            md_path = chapter_dir / "ch5.md"
            md_path.write_text("# 第5章\n\n图5.1 能源资源流向图\n", encoding="utf-8")
            doc = Document()
            BUILD_MOD._render_md(
                doc,
                str(md_path),
                {
                    "images": {
                        "图5.1 能源资源流向图": ["data/images/flow.png"]
                    },
                    "project_dir": str(root),
                },
            )
            self.assertAlmostEqual(self._picture_widths_cm(doc)[0], 12.0, places=1)


if __name__ == "__main__":
    unittest.main()
