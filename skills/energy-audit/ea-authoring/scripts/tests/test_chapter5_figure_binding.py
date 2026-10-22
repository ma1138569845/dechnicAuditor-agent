import base64
import hashlib
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT.parents[1] / "energy-audit-report" / "scripts" / "build_energy_audit_docx.py"
ASSERTS = ROOT.parents[1] / "energy-audit-report" / "scripts" / "ea_docx_asserts.py"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


BUILD_MOD = load_module("build_energy_audit_docx", BUILD)
ASSERTS_MOD = load_module("ea_docx_asserts_ch5", ASSERTS)


class BindingProject:
    def __init__(self, root: Path):
        self.root = root
        (root / "charts").mkdir()
        (root / "chapter_md").mkdir()
        chart = root / "charts" / "energy_flow.png"
        chart.write_bytes(PNG)
        digest = hashlib.sha256(chart.read_bytes()).hexdigest()
        caption = "图5.1 能源资源流向图"
        self.manifest = {
            "schema_version": 1,
            "batch_id": "test",
            "figures": [
                {
                    "figure_no": "5.1",
                    "caption": caption,
                    "src": "charts/energy_flow.png",
                    "kind": "energy_flow",
                    "energy_code": "",
                    "sha256": digest,
                }
            ],
        }
        (root / "chapter5_figures.json").write_text(
            json.dumps(self.manifest, ensure_ascii=False), encoding="utf-8"
        )
        (root / "chapter_md" / "ch5_import.md").write_text(
            f"# 第5章\n\n{caption}\n", encoding="utf-8"
        )
        (root / "report_images.json").write_text(
            json.dumps(
                {
                    "images": [
                        {
                            "caption": caption,
                            "src": "charts/energy_flow.png",
                            "sha256": digest,
                        }
                    ]
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )


class Chapter5AssemblyPreflightTests(unittest.TestCase):
    def test_matching_manifest_markdown_and_images_pass(self):
        with tempfile.TemporaryDirectory() as td:
            project = BindingProject(Path(td))
            result = BUILD_MOD._validate_chapter5_figures(str(project.root))
            self.assertEqual(result["manifest"]["batch_id"], "test")
            self.assertEqual(result["issues"], [])

    def test_missing_chart_file_degrades_to_placeholder(self):
        with tempfile.TemporaryDirectory() as td:
            project = BindingProject(Path(td))
            (project.root / "charts" / "energy_flow.png").unlink()
            result = BUILD_MOD._validate_chapter5_figures(str(project.root))
            self.assertEqual(result["issues"][0]["code"], "chart_file_missing")
            self.assertEqual(result["safe_images"], {})

    def test_caption_mismatch_degrades_to_placeholder(self):
        with tempfile.TemporaryDirectory() as td:
            project = BindingProject(Path(td))
            md = project.root / "chapter_md" / "ch5_import.md"
            md.write_text("# 第5章\n\n图5.1 错误图注\n", encoding="utf-8")
            result = BUILD_MOD._validate_chapter5_figures(str(project.root))
            self.assertIn(
                "caption_mismatch",
                [issue["code"] for issue in result["issues"]],
            )
            self.assertEqual(result["safe_images"], {})

    def test_report_image_mismatch_degrades_to_placeholder(self):
        with tempfile.TemporaryDirectory() as td:
            project = BindingProject(Path(td))
            path = project.root / "report_images.json"
            data = json.loads(path.read_text(encoding="utf-8"))
            data["images"][0]["src"] = "charts/other.png"
            path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            result = BUILD_MOD._validate_chapter5_figures(str(project.root))
            self.assertEqual(
                result["issues"][0]["code"], "report_images_mismatch"
            )
            self.assertEqual(result["safe_images"], {})

    def test_strict_policy_still_blocks_missing_chart(self):
        with tempfile.TemporaryDirectory() as td:
            project = BindingProject(Path(td))
            (project.root / "charts" / "energy_flow.png").unlink()
            with self.assertRaises(BUILD_MOD.FigureBindingError):
                BUILD_MOD._validate_chapter5_figures(
                    str(project.root), strict=True
                )

    def test_placeholder_is_visible_and_has_no_image(self):
        with tempfile.TemporaryDirectory() as td:
            project = BindingProject(Path(td))
            doc = Document()
            BUILD_MOD._figure(
                doc,
                "图5.1 能源资源流向图",
                {},
                str(project.root),
                figure_issues={
                    "图5.1 能源资源流向图": {
                        "reason": "未取得可用的本次图表",
                    }
                },
            )
            self.assertEqual(len(doc.paragraphs), 1)
            text = doc.paragraphs[0].text
            self.assertIn("图5.1 能源资源流向图", text)
            self.assertIn("暂缺", text)
            self.assertIn("未取得可用的本次图表", text)
            self.assertNotIn("<w:drawing", doc.paragraphs[0]._p.xml)


class Chapter5DocxAssertTests(unittest.TestCase):
    @staticmethod
    def para(text="", image=False):
        drawing = "<w:r><w:drawing/></w:r>" if image else ""
        return f"<w:p><w:r><w:t>{text}</w:t></w:r>{drawing}</w:p>"

    def run_assert(self, blocks):
        with mock.patch.object(ASSERTS_MOD, "_chapter_region", return_value=blocks):
            return ASSERTS_MOD.ch5_figures_ok("<w:body/>")

    def test_each_caption_has_exactly_one_preceding_image(self):
        ok, why = self.run_assert(
            [
                self.para(image=True),
                self.para("图5.1 能源资源流向图"),
                self.para(image=True),
                self.para("图5.2 2022年-2024年总用电量（单位：kWh）"),
            ]
        )
        self.assertTrue(ok, why)

    def test_caption_without_image_fails(self):
        ok, why = self.run_assert([self.para("图5.1 能源资源流向图")])
        self.assertFalse(ok)
        self.assertIn("没有图片", why)

    def test_visible_missing_placeholder_is_accepted(self):
        ok, why = self.run_assert(
            [
                self.para(
                    "图5.1 能源资源流向图（暂缺：未取得可用的本次图表）"
                )
            ]
        )
        self.assertTrue(ok, why)

    def test_non_continuous_numbers_fail(self):
        ok, why = self.run_assert(
            [
                self.para(image=True),
                self.para("图5.1 能源资源流向图"),
                self.para(image=True),
                self.para("图5.3 2022年能源费用占比"),
            ]
        )
        self.assertFalse(ok)
        self.assertIn("不连续", why)


if __name__ == "__main__":
    unittest.main()
