# -*- coding: utf-8 -*-
"""能源审计报告装配器（脚本化装配 B1）—— T1 骨架 + T2 章节 + T3 附录 + T4 页眉水印页脚

一次构建报告 docx 的固定版式 + 章节正文 + 附录 + 页眉/水印/页脚：
  封面（3空 + 单位22pt + 报告名26pt + 审计期间 + 8空 + 机构/日期 + 分页）
  三张信息表（能源审计机构 / 审计组 / 配合人员，数据源 data.json）
  目录页（'目  录' 标题无标题样式防自收录 + TOC 域 \\o "1-3"，其后不插分页）
  第 1~8 章渲染（chapter_md/chN.md → Heading 样式 + 正文 + 表格 + 图 + OMML 公式）
  附录渲染（chapter_md/appendix.md → '附录：'总页 + 附表 + 附录说明文字）

T2 章节渲染（md 元素映射，对齐 45 页终稿）：
  - H1 宋体15pt居中 / H2 宋体14pt / H3 宋体12pt（Heading 样式 + 直接格式）
  - 正文：1.5 倍行距、两端对齐、首行缩进 2 字符（firstLineChars=200）
  - 表格：Table Grid、12pt 居中、垂直居中、行高 1.01cm、首行加粗
  - 表题/子标题（**…**）：居中 12pt 加粗；图注：居中 12pt 不加粗
  - 项目符号：Wingdings U+F06C 实心圆点（numId=1 的 abstractNum lvl0 重定义）
  - 公式：md 占位 [FORMULAn] → assets/omml_formulas.json 注入 OMML（居中段）
  - 图片：report_images.json 清单（图注精确匹配 → 12cm 独立居中段 + 图注段）

T3 附录渲染特例：
  - '# 附录：'：H1 样式但 12pt 不加粗左对齐
  - '附录N：…' 清单行：1.5 行距、无缩进、无对齐覆盖
  - '附表X-Y …' 行：居中 12pt 加粗表题；其余同章节（正文/表格/说明段）

T4 页眉/水印/页脚（构建内直写，对齐 45 页终稿）：
  - 页眉：单位全称 + 两空格 + "能源审计报告"，右对齐宋体 10.5pt + 底边线（pBdr）
  - 水印：EAWatermark DrawingML（behindDoc=1、页面居中）——assets/header_template.xml
  - 页脚：— PAGE —（居中 10.5pt）——assets/footer_template.xml
  - settings.xml 写 updateFields（打开提示更新域）；模板单位名以 {{UNIT}} 占位注入。
  模板来源：烟台法院 45 页终稿（2026-09 验收版），随技能资产维护。

后续阶段（T5 收尾 PDF）在此骨架上叠加。

用法:
  python build_energy_audit_docx.py --project-dir <项目目录> [--out <docx路径>]
  数据源: <项目目录>/data.json + chapter_md/ + report_images.json
  默认输出: <项目目录>/output/_script_build/<单位全称>能源审计报告.docx

格式基准: ea-authoring/references/docx-ooxml-techniques.md
交付基准: 烟台经开区法院 45 页终稿（2026-09）
"""
import argparse
import json
import os
import re
import sys
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SONG = "宋体"
TNR = "Times New Roman"
HEAD_COLOR = RGBColor(0x1A, 0x1A, 0x1A)  # 与终稿标题色一致
W_NS_DECL = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
BULLET_CHAR = "\uf06c"  # Wingdings 实心圆点（对齐 45 页终稿）

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_ASSETS_DIR = os.path.normpath(os.path.join(_SCRIPT_DIR, "..", "assets"))

CHAPTERS = tuple(f"ch{n}" for n in range(1, 9))

_THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


def _ensure_rfonts(rpr):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in _THEME_ATTRS:  # 主题字体属性优先级高于显式值，必须清除
        rf.attrib.pop(qn(k), None)
    rf.set(qn("w:ascii"), TNR)
    rf.set(qn("w:hAnsi"), TNR)
    rf.set(qn("w:eastAsia"), SONG)
    rf.set(qn("w:cs"), TNR)
    return rf


def set_font(run, size=12, bold=False, color=None):
    """统一设置 run 字体：西文 TNR + 中文宋体。"""
    run.font.name = TNR
    run.font.size = Pt(size)
    run.font.bold = bold
    _ensure_rfonts(run._element.get_or_add_rPr())
    if color is not None:
        run.font.color.rgb = color


def _para(doc, text="", *, size=12, bold=False, align=None, line=None,
          space_before=None, space_after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if line is not None:
        pf.line_spacing = line
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if text:
        set_font(p.add_run(text), size, bold)
    return p


def _spacers(doc, n):
    for _ in range(n):
        _para(doc, "", line=1.5, space_before=0, space_after=0)


def _page_break(doc):
    p = doc.add_paragraph()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    p.add_run()._r.append(br)
    return p


def _add_toc_field(paragraph):
    """插入 TOC 域（\\o "1-3"），占位文本待 Word 打开/收尾脚本更新。"""
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); run._r.append(f1)
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" '
    run._r.append(instr)
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate"); run._r.append(f2)
    t = OxmlElement("w:t"); t.text = "（打开文档后目录将自动更新）"; run._r.append(t)
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end"); run._r.append(f3)
    set_font(run, 12)


def _add_table(doc, rows):
    """通用表格：12pt 居中、垂直居中、行高 1.01cm，首行加粗（对齐 45 页终稿）。"""
    n_cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n_cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, data in enumerate(rows):
        row = t.rows[ri]
        row.height = Cm(1.01)
        for ci in range(n_cols):
            cell = row.cells[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            set_font(p.add_run(data[ci] if ci < len(data) else ""), 12, ri == 0)
    return t


def _info_title(doc, text):
    return _para(doc, text, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)


def _build_cover(doc, unit, period, org, date_text):
    _spacers(doc, 3)
    for text, size, bold in ((unit, 22, True), ("能源审计报告", 26, True), (period, 14, False)):
        if text:
            _para(doc, text, size=size, bold=bold, align=WD_ALIGN_PARAGRAPH.CENTER,
                  line=1.5, space_before=0, space_after=0)
    _spacers(doc, 8)
    for text in (f"审计机构：{org}", date_text):
        if text:
            _para(doc, text, size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
                  line=1.5, space_before=0, space_after=0)
    _page_break(doc)


def _build_info_page(doc, institution, team, coop):
    _info_title(doc, "能源审计机构信息表")
    _add_table(doc, institution)
    _info_title(doc, "能源审计组人员名单")
    _add_table(doc, team)
    _info_title(doc, "能源审计配合人员名单")
    _add_table(doc, coop)
    doc.add_paragraph()
    _page_break(doc)


def _build_toc_page(doc):
    _para(doc, "目  录", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5)
    _add_toc_field(doc.add_paragraph())


def _setup_page(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(3.17)
    s.header_distance = Cm(1.5)
    s.footer_distance = Cm(1.75)


def _setup_styles(doc):
    """标题样式定义：宋体 15/14/12 加粗 + 大纲级别；Normal 宋体 12pt。"""
    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(12)
    _ensure_rfonts(normal.element.get_or_add_rPr())
    for lv, size in ((1, 15), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {lv}"]
        st.font.name = TNR
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = HEAD_COLOR
        _ensure_rfonts(st.element.get_or_add_rPr())
        pf = st.paragraph_format
        pf.keep_with_next = True
        pf.keep_together = True
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.7
        ppr = st.element.get_or_add_pPr()
        if ppr.find(qn("w:outlineLvl")) is None:
            ol = OxmlElement("w:outlineLvl")
            ol.set(qn("w:val"), str(lv - 1))
            ppr.append(ol)


def _setup_bullet_numbering(doc):
    """把 numId=1 引用的 abstractNum lvl0 重定义为 Wingdings 实心圆点（对齐终稿）。"""
    root = doc.part.numbering_part.element
    num = next((n for n in root.findall(qn("w:num")) if n.get(qn("w:numId")) == "1"), None)
    if num is None:
        print("WARN: numbering.xml 无 numId=1，项目符号不可用")
        return
    aid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    ab = next((a for a in root.findall(qn("w:abstractNum"))
               if a.get(qn("w:abstractNumId")) == aid), None)
    if ab is None:
        print("WARN: abstractNum 缺失（numId=1）")
        return
    old = next((l for l in ab.findall(qn("w:lvl")) if l.get(qn("w:ilvl")) == "0"), None)
    new = parse_xml(
        f'<w:lvl {W_NS_DECL} w:ilvl="0">'
        '<w:start w:val="1"/><w:numFmt w:val="bullet"/>'
        f'<w:lvlText w:val="{BULLET_CHAR}"/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="420" w:hanging="420"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr>'
        '</w:lvl>')
    if old is not None:
        ab.replace(old, new)
    else:
        ab.append(new)


def _set_update_fields(doc):
    """settings.xml 写 <w:updateFields w:val="true"/>（打开时提示更新域）。"""
    el = doc.settings.element
    if el.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        el.append(uf)


def _ensure_hf_parts(doc):
    """先创建页眉/页脚部件（内容由 _inject_header_footer 按模板 zip 级注入）。"""
    sec = doc.sections[0]
    for hf in (sec.header, sec.footer):
        hf.is_linked_to_previous = False


def _inject_header_footer(out_path, unit):
    """把 assets 模板（{{UNIT}} 替换）注入 headerN/footerN 部件（对齐 45 页终稿）。"""
    hdr_path = os.path.join(_ASSETS_DIR, "header_template.xml")
    ftr_path = os.path.join(_ASSETS_DIR, "footer_template.xml")
    if not (os.path.exists(hdr_path) and os.path.exists(ftr_path)):
        print("WARN: 页眉/页脚模板缺失，跳过注入")
        return
    with open(hdr_path, encoding="utf-8") as f:
        hdr = f.read().replace("{{UNIT}}", unit)
    with open(ftr_path, encoding="utf-8") as f:
        ftr = f.read()
    out_path = os.fspath(out_path)
    tmp = out_path + ".tmp"
    with zipfile.ZipFile(out_path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            fn = item.filename
            if re.fullmatch(r"word/header\d*\.xml", fn):
                dst.writestr(item, hdr.encode("utf-8"))
            elif re.fullmatch(r"word/footer\d*\.xml", fn):
                dst.writestr(item, ftr.encode("utf-8"))
            else:
                dst.writestr(item, src.read(fn))
    os.replace(tmp, out_path)


# ---------------------------------------------------------------- T2 章节渲染

def _body(doc, text):
    """正文自然段：两端对齐、1.5 行距、首行缩进 2 字符（firstLineChars=200）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Pt(24)
    ppr = p._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is not None:
        ind.set(qn("w:firstLineChars"), "200")
    _add_rich(p, text, 12)
    return p


def _add_rich(p, text, size=12):
    """按 **…** 拆分内联加粗。"""
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            set_font(p.add_run(seg[2:-2]), size, True)
        else:
            set_font(p.add_run(seg), size, False)


def _heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), {1: 15, 2: 14, 3: 12}[level], True)
    return p


def _apx_h1(doc, text):
    """附录总页标题：H1 样式但 12pt 不加粗左对齐（对齐终稿）。"""
    p = doc.add_paragraph(style="Heading 1")
    set_font(p.add_run(text), 12, False)
    return p


def _plain_line(doc, text):
    """附录清单行：1.5 行距、无缩进、无对齐覆盖（对齐终稿）。"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    _add_rich(p, text, 12)
    return p


def _bullet(doc, text):
    """列表项：numId=1（Wingdings 圆点）+ 1.5 行距；缩进由编号定义提供。"""
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId"); numid.set(qn("w:val"), "1")
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)
    p.paragraph_format.line_spacing = 1.5
    _add_rich(p, text, 12)
    return p


def _add_image_para(doc, path, width_cm=12.0):
    """图片独立段：居中、12cm 宽、无缩进（对齐规范；终稿缩进漂移不复制）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    r = p.add_run()
    r.add_picture(path, width=Cm(width_cm))
    return p


def _formula_para(doc, omml_xml):
    """公式段：居中、无缩进、间距 0/0、1.5 行距；内容为 OMML。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    p._p.append(parse_xml(omml_xml))
    return p


def _render_formula_line(doc, line, omml):
    """[FORMULAn] 占位行 → 前段 / 公式段 / 后段（对齐终稿三段式）。"""
    for k, seg in enumerate(re.split(r"\[FORMULA(\d+)\]", line)):
        if k % 2 == 0:
            if seg.strip():
                _body(doc, seg.strip())
        else:
            xml = omml.get(f"FORMULA{seg}")
            if xml:
                _formula_para(doc, xml)
            else:
                print(f"WARN: 公式库缺少 FORMULA{seg}")


def _figure(doc, caption, images, project_dir):
    """图注行 → 先插清单对应的图（可多张），再插居中图注段。"""
    srcs = images.get(caption) or []
    if not srcs:
        print(f"WARN: 图清单缺少: {caption}")
    for src in srcs:
        _add_image_para(doc, os.path.join(project_dir, src))
    _para(doc, caption, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)


def _load_omml():
    path = os.path.join(_ASSETS_DIR, "omml_formulas.json")
    if not os.path.exists(path):
        print(f"WARN: 公式库缺失: {path}")
        return {}
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _load_images(project_dir):
    path = os.path.join(project_dir, "report_images.json")
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    out = {}
    for e in data.get("images", []):
        cap = (e.get("caption") or "").strip()
        if cap and e.get("src"):
            out.setdefault(cap, []).append(e["src"])
    return out


def _render_md(doc, md_path, ctx, appendix=False):
    """行级 md 渲染器（内容层不变：只做元素映射，不做写作）。

    appendix=True 启用附录特例：'# 附录：' H1=12pt 不加粗左对齐；
    '附录N：…' 清单行=1.5 行距无缩进；'附表X-Y …' 行=居中加粗表题。
    """
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    omml = ctx.get("omml") or {}
    images = ctx.get("images") or {}
    project_dir = ctx["project_dir"]
    listing = False  # 附录总目录清单区（'附录：' 之后、首个 '## ' 之前）
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("### "):
            _heading(doc, line[4:].strip(), 3)
        elif line.startswith("## "):
            listing = False
            _heading(doc, line[3:].strip(), 2)
        elif line.startswith("# "):
            text = line[2:].strip()
            if appendix:
                _apx_h1(doc, text)
                listing = text.startswith("附录")
            else:
                _heading(doc, text, 1)
        elif line.startswith("- "):
            _bullet(doc, line[2:].strip())
        elif line.startswith("> "):
            _body(doc, line[2:].strip())
        elif line.startswith("|"):
            rows = []
            while True:
                cells = [c.strip() for c in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                if i >= len(lines):
                    break
                line = lines[i].strip()
                if not line.startswith("|"):
                    break
                i += 1
            if rows:
                _add_table(doc, rows)
        elif "[FORMULA" in line:
            _render_formula_line(doc, line, omml)
        elif line.startswith("**表") and line.endswith("**"):
            _para(doc, line[2:-2].strip(), size=12, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        elif re.match(r"^表\d+\s*[.\-–—]\s*\d+", line):
            # 裸表题（无 ** 包裹）：2026-09-20 新增，避免表题被当正文（两端对齐）导致 V3 P2
            _para(doc, line, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif appendix and re.match(r"^附表\d", line):
            _para(doc, line, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line.startswith("**") and line.endswith("**") and len(line) > 4:
            _para(doc, line[2:-2].strip(), size=12, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        elif re.match(r"^图\d+\.\d+", line):
            _figure(doc, line, images, project_dir)
        elif appendix and listing:
            _plain_line(doc, line)
        else:
            _body(doc, line)


def _render_chapters(doc, project_dir, ctx):
    md_dir = os.path.join(project_dir, "chapter_md")
    if not os.path.isdir(md_dir):
        print(f"WARN: 无 chapter_md 目录: {md_dir}")
        return
    for name in CHAPTERS:
        path = None
        for cand in (f"{name}_import.md", f"{name}.md"):
            p = os.path.join(md_dir, cand)
            if os.path.exists(p):
                path = p
                break
        if path is None:
            print(f"WARN: 缺章节 md: {name}")
            continue
        _render_md(doc, path, ctx)


def _render_appendix(doc, project_dir, ctx):
    path = os.path.join(project_dir, "chapter_md", "appendix.md")
    if not os.path.exists(path):
        print(f"WARN: 缺附录 md: {path}")
        return
    _render_md(doc, path, ctx, appendix=True)


def _info_rows(base, data):
    institution = [
        ["机构名称", base.get("audit_org_name") or ""],
        ["地址", base.get("audit_org_address") or ""],
        ["负责人", base.get("audit_org_contact") or ""],
        ["联系方式", base.get("audit_org_phone") or ""],
    ]
    team_rows = [
        [m.get("role", ""), m.get("name", ""), m.get("education", ""),
         m.get("certification", ""), m.get("major", "")]
        for m in data.get("audit_team", [])
    ]
    coop_rows = [
        [c.get("role", ""), c.get("dept", ""), c.get("name", ""),
         c.get("gender", ""), c.get("position", "")]
        for c in data.get("cooperation", [])
    ]
    team = [["组内职务", "姓名", "学历", "所获资质", "专业"]] + (team_rows or [["【待补充】"] * 5])
    coop = [["组内职务", "部门", "姓名", "性别", "职务"]] + (coop_rows or [["【待补充】"] * 5])
    return institution, team, coop


def build(project_dir, out_path=None):
    with open(os.path.join(project_dir, "data.json"), encoding="utf-8") as f:
        data = json.load(f)
    base = data.get("base", {})
    unit = base.get("unit_name") or "××单位"
    org = base.get("audit_org_name") or ""
    date_text = base.get("report_date") or ""
    years = sorted({y.get("year") for y in data.get("energy_yearly", []) if y.get("year")})
    period = f"（审计期间：{years[0]}年—{years[-1]}年）" if years else ""

    if not out_path:
        out_path = os.path.join(project_dir, "output", "_script_build",
                                f"{unit}能源审计报告.docx")

    doc = Document()
    _setup_page(doc)
    _setup_styles(doc)
    _setup_bullet_numbering(doc)
    _set_update_fields(doc)
    _ensure_hf_parts(doc)
    _build_cover(doc, unit, period, org, date_text)
    _build_info_page(doc, *_info_rows(base, data))
    _build_toc_page(doc)
    ctx = {
        "project_dir": project_dir,
        "omml": _load_omml(),
        "images": _load_images(project_dir),
    }
    _render_chapters(doc, project_dir, ctx)
    _render_appendix(doc, project_dir, ctx)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    _inject_header_footer(out_path, unit)
    return out_path, doc


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="能源审计报告装配器（T1~T4：封面+信息表+目录+第1~8章+附录+页眉水印页脚）")
    ap.add_argument("--project-dir", required=True, help="项目目录（含 data.json / chapter_md/）")
    ap.add_argument("--out", help="输出 docx 路径（默认 output/_script_build/）")
    a = ap.parse_args(argv)
    out, doc = build(a.project_dir, a.out)
    om = len(doc.element.body.findall(".//" + qn("m:oMath")))
    print(f"OK -> {out} | paras={len(doc.paragraphs)} tables={len(doc.tables)} oMath={om}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
