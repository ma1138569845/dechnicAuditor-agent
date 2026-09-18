# -*- coding: utf-8 -*-
"""装配器测试（T1 骨架 + T2 章节 + T3 附录 + T4 页眉水印页脚）：
合成数据 → build_energy_audit_docx.py → docx 结构断言。

覆盖 T1：封面结构（空行数/字号/分页符）、三张信息表（形状/表头/样式）、
目录页（无标题样式防自收录 + TOC 域）、页面设置、标题样式定义。
覆盖 T2：章节 H1/H2/H3 样式、正文缩进/行距/对齐、项目符号（numId=1 + Wingdings）、
章表 + 表题、OMML 公式注入（[FORMULAn] 三段式）、图片（清单匹配）+ 图注、子标题。
覆盖 T3：'附录：' H1（12pt 不加粗左对齐）、附录清单行（1.5 行距无缩进）、
H2 附录标题、附表题（居中加粗）、附录表格。
覆盖 T4：页眉（单位名+报告名+pBdr）、水印（EAWatermark/behindDoc、无 textpath）、
页脚（— PAGE —）、settings updateFields（zip 级校验）。
覆盖 T5：收尾器（finalize_energy_audit_pdf.py）--help 冒烟（COM 路径不在单元测试内）。
"""
from __future__ import annotations

import base64
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

REPO = Path(__file__).resolve().parent.parent.parent
SCRIPT = (REPO / "skills" / "energy-audit" / "energy-audit-report" / "scripts"
          / "build_energy_audit_docx.py")

UNIT = "测试市第一中学"
ORG = "测试德诚科技有限公司"
# 1x1 透明 PNG（最小合法图片，供图片清单用例）
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def _make_project(root: Path) -> Path:
    proj = root / "proj"
    proj.mkdir(parents=True)
    data = {
        "base": {
            "unit_name": UNIT,
            "audit_org_name": ORG,
            "audit_org_address": "测试省测试市测试路1号",
            "audit_org_contact": "王五",
            "audit_org_phone": "13800000000",
            "report_date": "2026年1月",
        },
        "energy_yearly": [{"year": 2023}, {"year": 2024}, {"year": 2025}],
        "audit_team": [
            {"role": "审计负责人", "name": "张三", "education": "硕士",
             "certification": "高级工程师", "major": "暖通空调"},
            {"role": "成员", "name": "赵六", "education": "本科",
             "certification": "工程师", "major": "电气自动化"},
        ],
        "cooperation": [
            {"role": "组长", "dept": "学校", "name": "李四", "gender": "女", "position": "主任"},
        ],
    }
    (proj / "data.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    # 章节 md（覆盖 T2 全元素）
    chap = proj / "chapter_md"
    chap.mkdir()
    (chap / "ch1.md").write_text(
        "# 第1章 测试章\n\n"
        "## 1.1 测试节\n\n"
        "### 1.1.1 测试小节\n\n"
        "这是正文段落，包含内联**加粗**文字。\n\n"
        "- 列表项甲；\n"
        "- 列表项乙。\n\n"
        "**测试子标题**\n\n"
        "**表1.1 测试数据表**\n\n"
        "| 项目 | 数值 |\n"
        "| --- | --- |\n"
        "| 甲 | 1 |\n\n"
        "单位面积值按式[FORMULA1]计算（式中：E 为综合能耗）。\n\n"
        "另见式[FORMULA9]的说明。\n\n"
        "图1.1 测试示意图\n\n"
        "尾段正文。\n",
        encoding="utf-8")

    # 附录 md（覆盖 T3 全元素）
    (chap / "appendix.md").write_text(
        "# 附录：\n\n"
        "附录1：测试附录\n\n"
        "## 附录1：测试附录\n\n"
        "附表1-1 测试附表\n\n"
        "| 甲 | 乙 |\n"
        "| --- | --- |\n"
        "| 1 | 2 |\n\n"
        "注：测试注记。\n",
        encoding="utf-8")

    # 图清单 + 最小图片
    img_dir = proj / "data" / "images"
    img_dir.mkdir(parents=True)
    (img_dir / "t.png").write_bytes(TINY_PNG)
    (proj / "report_images.json").write_text(
        json.dumps({"images": [{"caption": "图1.1 测试示意图",
                                "src": "data/images/t.png"}]}, ensure_ascii=False),
        encoding="utf-8")
    return proj


class _BuiltDoc:
    """Document 包装：属性委托给 Document，额外暴露 .out（产物路径，zip 级断言用）。"""

    def __init__(self, doc, out):
        object.__setattr__(self, "_doc", doc)
        self.out = out

    def __getattr__(self, name):
        return getattr(self._doc, name)


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    proj = _make_project(tmp_path_factory.mktemp("ea_build"))
    out = proj / "out" / "报告.docx"
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--project-dir", str(proj), "--out", str(out)],
        capture_output=True, text=True,
    )
    assert r.returncode == 0, f"stdout={r.stdout}\nstderr={r.stderr}"
    assert out.exists()
    built_doc = _BuiltDoc(Document(str(out)), out)
    built_doc.stdout = r.stdout
    return built_doc


# ---------------------------------------------------------------- T1

def test_cover_structure(built):
    paras = built.paragraphs
    # 3 空行 + 单位(22pt bold) + 报告名(26pt bold) + 期间(14pt)
    assert all(not p.text.strip() for p in paras[0:3])
    p_unit = paras[3]
    assert p_unit.text == UNIT
    assert p_unit.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p_unit.runs[0].font.size.pt == 22 and p_unit.runs[0].font.bold is True
    p_title = paras[4]
    assert p_title.text == "能源审计报告"
    assert p_title.runs[0].font.size.pt == 26 and p_title.runs[0].font.bold is True
    p_period = paras[5]
    assert p_period.text == "（审计期间：2023年—2025年）"
    assert p_period.runs[0].font.size.pt == 14 and not p_period.runs[0].font.bold
    # 8 空行 + 机构 + 日期
    assert all(not p.text.strip() for p in paras[6:14])
    assert paras[14].text == f"审计机构：{ORG}"
    assert paras[15].text == "2026年1月"
    # 分页符
    brs = [b.get(qn("w:type")) for b in paras[16]._p.findall(".//" + qn("w:br"))]
    assert brs == ["page"]
    # 封面段落字体：中文宋体
    rf = paras[3].runs[0]._element.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert rf.get(qn("w:eastAsia")) == "宋体"
    assert rf.get(qn("w:ascii")) == "Times New Roman"


def test_info_tables(built):
    assert len(built.tables) == 5  # 3 信息表 + 第1章表 + 附录附表
    t0, t1, t2 = built.tables[0], built.tables[1], built.tables[2]
    assert (len(t0.rows), len(t0.columns)) == (4, 2)
    assert (len(t1.rows), len(t1.columns)) == (3, 5)  # 表头 + 2 成员
    assert (len(t2.rows), len(t2.columns)) == (2, 5)
    assert t0.rows[0].cells[0].text == "机构名称"
    assert t0.rows[0].cells[1].text == ORG
    assert t0.rows[1].cells[0].text == "地址"
    assert t1.rows[0].cells[3].text == "所获资质"
    assert t1.rows[1].cells[1].text == "张三"
    assert t2.rows[1].cells[2].text == "李四"
    # 首行加粗、内容 12pt、单元格居中 + 垂直居中
    c0 = t0.rows[0].cells[0].paragraphs[0]
    assert c0.runs[0].font.bold is True and c0.runs[0].font.size.pt == 12
    assert c0.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert t0.rows[1].cells[0].paragraphs[0].runs[0].font.bold is not True
    assert t0.rows[0].cells[0].vertical_alignment is not None
    # 行高 1.01cm ≈ 572 twips
    tr_h = t0.rows[0]._tr.find(qn("w:trPr")).find(qn("w:trHeight"))
    assert int(tr_h.get(qn("w:val"))) in (571, 572, 573)
    # 表格居中
    assert "center" in t0._tbl.find(qn("w:tblPr")).find(qn("w:jc")).get(qn("w:val"))


def test_info_titles_and_order(built):
    paras = built.paragraphs
    idx = {p.text: i for i, p in enumerate(paras)}
    for title in ("能源审计机构信息表", "能源审计组人员名单", "能源审计配合人员名单"):
        i = idx[title]
        p = paras[i]
        assert p.style.name == "Normal"  # 不带标题样式（防 TOC 自收录）
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert p.runs[0].font.size.pt == 12 and p.runs[0].font.bold is True
    assert idx["能源审计机构信息表"] < idx["能源审计组人员名单"] < idx["能源审计配合人员名单"]
    # 三个标题都排在目录之前
    assert idx["能源审计配合人员名单"] < idx["目  录"]


def test_toc_page_and_field(built):
    paras = built.paragraphs
    toc_title = next(p for p in paras if p.text == "目  录")
    assert toc_title.style.name == "Normal"  # 防自收录：不得带 Heading 样式
    assert toc_title.runs[0].font.size.pt == 15 and toc_title.runs[0].font.bold is True
    assert toc_title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    xml = built.element.body.xml
    assert 'TOC \\o "1-3"' in xml
    # 目录后不得有分页符（第1章随目录后顺延，与 45 页终稿一致）
    toc_i = paras.index(toc_title)
    for p in paras[toc_i:]:
        brs = [b.get(qn("w:type")) for b in p._p.findall(".//" + qn("w:br"))]
        assert "page" not in brs


def test_page_setup_and_heading_styles(built):
    s = built.sections[0]
    assert (round(s.page_width.cm, 2), round(s.page_height.cm, 2)) == (21.0, 29.7)
    assert round(s.top_margin.cm, 2) == 2.54 == round(s.bottom_margin.cm, 2)
    assert round(s.left_margin.cm, 2) == 3.17 == round(s.right_margin.cm, 2)
    assert round(s.header_distance.cm, 2) == 1.5
    assert round(s.footer_distance.cm, 2) == 1.75
    for lv, size in ((1, 15), (2, 14), (3, 12)):
        st = built.styles[f"Heading {lv}"]
        assert st.font.size.pt == size and st.font.bold is True
        assert "宋体" in st.element.xml
        assert "outlineLvl" in st.element.xml


# ---------------------------------------------------------------- T2

def test_chapter_headings(built):
    idx = {p.text: p for p in built.paragraphs}
    h1 = idx["第1章 测试章"]
    assert h1.style.name == "Heading 1"
    assert h1.alignment == WD_ALIGN_PARAGRAPH.CENTER  # H1 居中
    assert h1.runs[0].font.size.pt == 15 and h1.runs[0].font.bold is True
    h2 = idx["1.1 测试节"]
    assert h2.style.name == "Heading 2"
    assert h2.runs[0].font.size.pt == 14
    h3 = idx["1.1.1 测试小节"]
    assert h3.style.name == "Heading 3"
    assert h3.runs[0].font.size.pt == 12


def test_body_para_format(built):
    p = next(p for p in built.paragraphs if p.text.startswith("这是正文段落"))
    assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert p.paragraph_format.line_spacing == 1.5
    ind = p._p.find(qn("w:pPr")).find(qn("w:ind"))
    assert ind.get(qn("w:firstLineChars")) == "200"
    assert ind.get(qn("w:firstLine")) == "480"
    # 内联加粗 run
    assert any(r.font.bold is True and r.text == "加粗" for r in p.runs)


def test_bullets(built):
    ps = [p for p in built.paragraphs if p.text.startswith("列表项")]
    assert len(ps) == 2
    for p in ps:
        assert p.text.startswith("列表项")  # 不含 "- " 前缀
        assert p.paragraph_format.line_spacing == 1.5
        numpr = p._p.find(qn("w:pPr")).find(qn("w:numPr"))
        assert numpr is not None
        assert numpr.find(qn("w:numId")).get(qn("w:val")) == "1"
        assert numpr.find(qn("w:ilvl")).get(qn("w:val")) == "0"
    # 编号定义已重定义为 Wingdings 圆点（兼容 PUA 字符被序列化为字符引用的情况）
    nx = built.part.numbering_part.element.xml
    assert ("\uf06c" in nx or "f06c" in nx.lower()) and "Wingdings" in nx


def test_chapter_table_and_caption(built):
    assert len(built.tables) == 5
    t = built.tables[3]
    assert (len(t.rows), len(t.columns)) == (2, 2)  # 表头 + 1 数据行（分隔行不计）
    assert t.rows[0].cells[0].text == "项目" and t.rows[0].cells[1].text == "数值"
    assert t.rows[0].cells[0].paragraphs[0].runs[0].font.bold is True
    assert t.rows[1].cells[0].text == "甲"
    cap = next(p for p in built.paragraphs if p.text == "表1.1 测试数据表")
    assert cap.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert cap.runs[0].font.size.pt == 12 and cap.runs[0].font.bold is True


def test_subtitle_line(built):
    p = next(p for p in built.paragraphs if p.text == "测试子标题")
    assert p.style.name == "Normal"
    assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p.runs[0].font.size.pt == 12 and p.runs[0].font.bold is True


def test_formula_injection(built):
    body = built.element.body
    assert len(body.findall(".//" + qn("m:oMath"))) == 1  # [FORMULA1] 已注入
    fp = next(p for p in built.paragraphs
              if p._p.findall(".//" + qn("m:oMath")))
    assert fp.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert fp.paragraph_format.line_spacing == 1.5
    # 三段式：前段「…按式」/ 公式段 / 后段「计算（式中…）」
    before = next(p for p in built.paragraphs if p.text.endswith("按式"))
    after = next(p for p in built.paragraphs if p.text.startswith("计算（式中"))
    assert before._p.find(qn("w:pPr")).find(qn("w:ind")) is not None
    assert after._p.find(qn("w:pPr")).find(qn("w:ind")) is not None


def test_figure_and_caption(built):
    paras = built.paragraphs  # 同一列表内取对象（python-docx 每次访问重建包装对象）
    body = built.element.body
    drawings = body.findall(".//" + qn("w:drawing"))
    assert len(drawings) == 1  # 清单图已插入
    img_p = next(p for p in paras if p._p.findall(".//" + qn("w:drawing")))
    assert img_p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    ppr = img_p._p.find(qn("w:pPr"))
    ind = ppr.find(qn("w:ind")) if ppr is not None else None
    assert ind is None or ind.get(qn("w:firstLineChars")) is None  # 图片段无缩进
    cap = next(p for p in paras if p.text == "图1.1 测试示意图")
    assert cap.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert cap.runs[0].font.size.pt == 12 and not cap.runs[0].font.bold
    # 图注排在图片段之后
    assert paras.index(img_p) < paras.index(cap)


# ---------------------------------------------------------------- T3

def test_appendix(built):
    paras = built.paragraphs
    # '附录：' H1 样式但 12pt 不加粗左对齐
    h1 = next(p for p in paras if p.text == "附录：")
    assert h1.style.name == "Heading 1"
    assert h1.runs[0].font.size.pt == 12 and not h1.runs[0].font.bold
    assert h1.alignment is None
    # '附录1：测试附录' 出现两次：清单行（Normal/无缩进）+ H2 标题
    occ = [p for p in paras if p.text == "附录1：测试附录"]
    assert len(occ) == 2
    lst, h2 = occ
    assert lst.style.name == "Normal"
    assert lst.paragraph_format.line_spacing == 1.5
    assert lst.alignment is None
    assert lst._p.find(qn("w:pPr")).find(qn("w:ind")) is None  # 无缩进
    assert h2.style.name == "Heading 2"
    # 附表题：居中加粗 12pt
    cap = next(p for p in paras if p.text == "附表1-1 测试附表")
    assert cap.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert cap.runs[0].font.bold is True and cap.runs[0].font.size.pt == 12
    # 附录表格在章表之后（tables: 3 信息 + 1 章 + 1 附录）
    assert len(built.tables) == 5
    t = built.tables[4]
    assert t.rows[0].cells[0].text == "甲" and t.rows[1].cells[1].text == "2"


# ---------------------------------------------------------------- T4

def test_header_footer_watermark_and_updatefields(built):
    with zipfile.ZipFile(built.out) as z:
        names = z.namelist()
        hdr = next(n for n in names if re.fullmatch(r"word/header\d*\.xml", n))
        ftr = next(n for n in names if re.fullmatch(r"word/footer\d*\.xml", n))
        h = z.read(hdr).decode("utf-8")
        f = z.read(ftr).decode("utf-8")
        s = z.read("word/settings.xml").decode("utf-8")
    # 页眉：单位名 + 报告名 + 底边线 + 水印（DrawingML，禁 textpath）
    assert UNIT in h and "能源审计报告" in h
    assert "pBdr" in h
    assert "EAWatermark" in h and 'behindDoc="1"' in h
    assert "textpath" not in h
    # 页脚：— PAGE — 域
    assert "PAGE" in f and "—" in f and "MERGEORMAT" not in f
    # updateFields
    assert "updateFields" in s


# ---------------------------------------------------------------- T5

def test_finalize_help():
    script = (REPO / "skills" / "energy-audit" / "energy-audit-report" / "scripts"
              / "finalize_energy_audit_pdf.py")
    r = subprocess.run([sys.executable, str(script), "--help"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert "usage" in r.stdout.lower()


# ---------------------------------------------------------------- T7

def test_missing_formula_warns(built):
    """公式库缺键：只告警不崩溃（[FORMULA9] 不在库中）。"""
    assert "WARN: 公式库缺少 FORMULA9" in built.stdout


def test_formula_variant_key_renders():
    """任意公式键可注入渲染（验证其他机构类型补入变体公式的机制）。"""
    import importlib.util
    spec = importlib.util.spec_from_file_location("build_ea", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    d = Document()
    mod._render_formula_line(
        d, "按式[FORMULA6]计算。", {"FORMULA6": mod._load_omml()["FORMULA1"]})
    assert len(d.paragraphs) == 3  # 前段 + 公式段 + 后段
    assert len(d.element.body.findall(".//" + qn("m:oMath"))) == 1
